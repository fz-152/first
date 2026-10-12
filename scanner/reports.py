"""Экспорт отчётов: PDF и Excel (ТЗ §4.2 «Экспорт»).

PDF формируется через ReportLab с поддержкой кириллицы (встроенный шрифт
DejaVuSans при наличии, иначе — Helvetica с транслитерацией).
"""
from __future__ import annotations

from pathlib import Path

from .checklists import ALL_BLOCKS, RULES_BY_ID, CRITICAL, MAJOR, MINOR
from .models import ScanResult

SEV_LABEL = {CRITICAL: "Критично (красный)", MAJOR: "Существенно (жёлтый)", MINOR: "Незначительно"}


# ------------------------------------------------------------------ Excel
def export_excel(res: ScanResult, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Резюме"
    fill_red = PatternFill("solid", fgColor="FFC7CE")
    fill_yel = PatternFill("solid", fgColor="FFEB9C")
    fill_grn = PatternFill("solid", fgColor="C6EFCE")
    light = _traffic(res.compliance_index)
    ws["A1"] = "Отчёт о соответствии сайта нормативным требованиям РФ и СПб"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Сайт", res.site),
        ("Дата начала", res.started_at), ("Дата окончания", res.finished_at),
        ("Страниц просканировано", res.pages_scanned),
        ("Индекс соответствия, %", res.compliance_index),
        ("Индикатор", {"green": "ЗЕЛЁНЫЙ", "yellow": "ЖЁЛТЫЙ", "red": "КРАСНЫЙ"}[light]),
        ("Критических нарушений", sum(1 for v in res.violations if v.severity == CRITICAL)),
        ("Существенных нарушений", sum(1 for v in res.violations if v.severity == MAJOR)),
        ("Незначительных замечаний", sum(1 for v in res.violations if v.severity == MINOR)),
    ]
    for i, (k, val) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=val)
        if k == "Индикатор":
            c.fill = {"green": fill_grn, "yellow": fill_yel, "red": fill_red}[light]
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60

    # лист нарушений, сгруппированных по блокам-законам
    ws2 = wb.create_sheet("Нарушения")
    headers = ["Правило", "Блок / закон", "Нарушение", "Описание", "URL страницы",
               "Угроза", "Ссылка на пункт закона"]
    for j, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=j, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    r = 2
    for block in ALL_BLOCKS:
        vs = [v for v in res.violations if v.block == block]
        if not vs:
            continue
        ws2.cell(row=r, column=1, value=block).font = Font(bold=True)
        r += 1
        for v in vs:
            vals = [v.rule_id, v.block, v.title, v.message, v.page_url,
                    SEV_LABEL[v.severity], v.law]
            for j, val in enumerate(vals, 1):
                cell = ws2.cell(row=r, column=j, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if j == 6:
                    cell.fill = {"critical": fill_red, "major": fill_yel, "minor": fill_grn}[v.severity]
            r += 1
        r += 1
    widths = (10, 34, 34, 60, 44, 18, 44)
    for j, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "A2"

    # лист пройденных проверок
    ws3 = wb.create_sheet("Пройдено")
    ws3.cell(row=1, column=1, value="Правило").font = Font(bold=True)
    ws3.cell(row=1, column=2, value="Проверка").font = Font(bold=True)
    for i, rid in enumerate(sorted(res.passed_rules), 2):
        rule = RULES_BY_ID.get(rid)
        if rule:
            ws3.cell(row=i, column=1, value=rid)
            ws3.cell(row=i, column=2, value=f"{rule.title} — {rule.law}")
    ws3.column_dimensions["A"].width = 12
    ws3.column_dimensions["B"].width = 90

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ------------------------------------------------------------------ PDF
def _find_cyrillic_font() -> str | None:
    for cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/TTF/DejaVuSans.ttf"):
        if Path(cand).exists():
            return cand
    import glob
    g = glob.glob("/usr/share/fonts/**/DejaVuSans.ttf", recursive=True)
    return g[0] if g else None


def _translit(s: str) -> str:
    table = {"а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "zh",
             "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
             "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts",
             "ч": "ch", "ш": "sh", "щ": "shch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
             "я": "ya", "і": "i", "ѐ": "e"}
    out = []
    for ch in s.lower():
        out.append(table.get(ch, ch))
    txt = "".join(out)
    return txt.upper() if s[:1].isupper() else txt


def _t(s: str, cyrillic: bool) -> str:
    return s if cyrillic else _translit(s)


def export_pdf(res: ScanResult, path: str | Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    font = _find_cyrillic_font()
    cyr = False
    base, bold = "Helvetica", "Helvetica-Bold"
    if font:
        pdfmetrics.registerFont(TTFont("DVS", font))
        try:
            pdfmetrics.registerFont(TTFont("DVSB", font))
        except Exception:
            pass
        base = bold = "DVS"
        cyr = True
    else:
        # ReportLab не умеет кириллицу в стандартных шрифтах — транслитерируем
        from reportlab.pdfbase.pdfdoc import PDFDocument  # noqa: F401 (тихо)
        def _cvt(text, *a, **kw):
            return _translit(text)
        base = bold = "Helvetica"

    styles = getSampleStyleSheet()
    title_st = ParagraphStyle("ttl", parent=styles["Title"], fontName=bold, fontSize=15)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName=bold, fontSize=12)
    body = ParagraphStyle("bd", parent=styles["BodyText"], fontName=base, fontSize=9, leading=12)
    small = ParagraphStyle("sm", parent=body, fontSize=8, textColor=colors.HexColor("#555555"))

    def P(text):
        return Paragraph(_t(text, cyr), body)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=_translit("Отчёт соответствия сайта"))
    story = [Paragraph(_t("Отчёт о соответствии сайта образовательной организации "
                          "нормативным требованиям РФ и Санкт-Петербурга", cyr), title_st),
             Spacer(1, 6 * mm)]
    light = _traffic(res.compliance_index)
    color = {"green": colors.green, "yellow": colors.orange, "red": colors.red}[light]
    meta = Table([[_t("Сайт", cyr), _t(res.site, cyr)],
                  [_t("Период сканирования", cyr), _t(f"{res.started_at} — {res.finished_at}", cyr)],
                  [_t("Страниц просканировано", cyr), str(res.pages_scanned)],
                  [_t("Индекс соответствия", cyr), f"{res.compliance_index} %"],
                  [_t("Индикатор", cyr), _t({"green": "ЗЕЛЁНЫЙ", "yellow": "ЖЁЛТЫЙ",
                                             "red": "КРАСНЫЙ"}[light], cyr)]],
                 colWidths=[55 * mm, 105 * mm])
    meta.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), base),
                              ("FONTSIZE", (0, 0), (-1, -1), 9),
                              ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                              ("TEXTCOLOR", (1, 3), (1, 4), color),
                              ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [meta, Spacer(1, 5 * mm)]

    for note in res.notes:
        story.append(Paragraph(_t(note, cyr), small))
    story.append(Spacer(1, 3 * mm))

    sev_order = [CRITICAL, MAJOR, MINOR]
    for block in ALL_BLOCKS:
        vs = [v for v in res.violations if v.block == block]
        if not vs:
            continue
        story.append(Paragraph(_t(block, cyr), h2))
        data = [[_t("№", cyr), _t("Нарушение", cyr), _t("Страница", cyr),
                 _t("Основание (закон)", cyr), _t("Угроза", cyr)]]
        row_colors = []
        for i, v in enumerate(vs, 1):
            data.append([v.rule_id, _t(v.title + ". " + v.message, cyr),
                         _t(v.page_url or "—", cyr), _t(v.law, cyr),
                         _t({"critical": "КРИТИЧНО", "major": "СУЩЕСТВЕННО",
                             "minor": "незначительно"}[v.severity], cyr)])
            if v.severity == CRITICAL:
                row_colors.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#fde8e8")))
            elif v.severity == MAJOR:
                row_colors.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#fff7d6")))
        tbl = Table(data, colWidths=[12 * mm, 62 * mm, 40 * mm, 36 * mm, 20 * mm], repeatRows=1)
        style = [("FONTNAME", (0, 0), (-1, -1), base), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                 ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce6f1")),
                 ("FONTNAME", (0, 0), (-1, 0), bold)] + row_colors
        tbl.setStyle(TableStyle(style))
        story += [tbl, Spacer(1, 4 * mm)]

    passed = sorted(set(res.passed_rules))
    if passed:
        story.append(Paragraph(_t("Пройденные проверки", cyr), h2))
        pd = [[_t("Правило", cyr), _t("Проверка / основание", cyr)]]
        for rid in passed:
            r = RULES_BY_ID.get(rid)
            if r:
                pd.append([rid, _t(f"{r.title} — {r.law}", cyr)])
        pt = Table(pd, colWidths=[22 * mm, 148 * mm], repeatRows=1)
        pt.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), base), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce6f1"))]))
        story.append(pt)

    doc.build(story)
    return path


def _traffic(index: float) -> str:
    from .engine import traffic_light
    return traffic_light(index)
